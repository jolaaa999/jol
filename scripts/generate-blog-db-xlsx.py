# -*- coding: utf-8 -*-
"""Generate JOL blog database design Excel (AgentScope-style)."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parents[1] / "docs" / "jol_博客数据表.xlsx"

TITLE_FONT = Font(name="微软雅黑", size=16, bold=True, color="000000")
SECTION_FONT = Font(name="微软雅黑", size=12, bold=True)
HEADER_FONT = Font(name="微软雅黑", size=11)
BODY_FONT = Font(name="微软雅黑", size=11)
NOTE_FONT = Font(name="微软雅黑", size=11)
HEADER_FILL = PatternFill("solid", fgColor="BFBFBF")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
WRAP = Alignment(wrap_text=True, vertical="center")
LEFT = Alignment(wrap_text=True, vertical="center", horizontal="left")


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def style_header_row(ws, row, cols=5):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP
        cell.border = THIN


def style_body_row(ws, row, cols=5):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.font = BODY_FONT
        cell.alignment = WRAP
        cell.border = THIN


def main() -> None:
    wb = Workbook()

    # ── 总览 ──
    ws = wb.active
    ws.title = "总览"
    set_widths(ws, [22, 18, 36, 18, 42])

    ws.merge_cells("A1:E2")
    ws["A1"] = "JOL 博客数据库表设计（MySQL · articles）"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(vertical="center", wrap_text=True)

    ws.merge_cells("A4:E4")
    ws["A4"] = "表清单"
    ws["A4"].font = SECTION_FONT

    for i, h in enumerate(["表名", "中文名", "说明", "所属模块", "备注"], 1):
        ws.cell(5, i, h)
    style_header_row(ws, 5)

    for i, v in enumerate(
        [
            "articles",
            "文章",
            "博客文章正文与元数据（有感 / 诗文）",
            "内容",
            "主键 id；按 category + created_at 索引",
        ],
        1,
    ):
        ws.cell(6, i, v)
    style_body_row(ws, 6)

    ws.merge_cells("A8:E8")
    ws["A8"] = "设计约定"
    ws["A8"].font = SECTION_FONT

    notes = [
        "1. 当前博客核心实体仅 articles 一张表；分类用 category 字段区分「有感」「诗文」，对应 API /api/posts 与 /api/poetry。",
        "2. tags 使用 MySQL JSON 数组存储（如 [\"design\",\"vue\"]）；无标签时存 [] 或 NULL。",
        "3. 写接口（POST/PUT/DELETE）需 ADMIN_TOKEN；读接口在未配置 MYSQL_DSN 时回退内置 mock。",
        "4. id 为业务主键（varchar），新建时可省略，由服务端按分类前缀生成（r- / p- + slug + 时分秒）。",
        "5. 约定：MySQL 8.0 / InnoDB / utf8mb4；字段以 backend/sql/schema.sql 为准。",
    ]
    for i, text in enumerate(notes):
        r = 9 + i
        ws.merge_cells(f"A{r}:E{r}")
        ws[f"A{r}"] = text
        ws[f"A{r}"].font = NOTE_FONT
        ws[f"A{r}"].alignment = LEFT
        ws.row_dimensions[r].height = 28

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18

    # ── 内容 ──
    ws2 = wb.create_sheet("内容")
    set_widths(ws2, [18, 22, 28, 42, 48])

    ws2.merge_cells("A1:E2")
    ws2["A1"] = "博客文章与分类内容"
    ws2["A1"].font = TITLE_FONT
    ws2["A1"].alignment = Alignment(vertical="center", wrap_text=True)

    ws2.merge_cells("A4:E4")
    ws2["A4"] = "文章表（articles）"
    ws2["A4"].font = SECTION_FONT

    for i, h in enumerate(["字段名", "类型", "说明", "约束", "备注"], 1):
        ws2.cell(5, i, h)
    style_header_row(ws2, 5)

    fields = [
        ("id", "varchar(64)", "文章业务主键", "PRIMARY KEY, NOT NULL", "示例：r-001、p-001；新建可省略，服务端自动生成"),
        ("title", "varchar(255)", "标题", "NOT NULL", None),
        ("category", "varchar(32)", "分类", "NOT NULL", "有感 | 诗文；有感走 /api/posts，诗文走 /api/poetry"),
        ("content", "mediumtext", "正文（Markdown）", "NOT NULL", "前端 marked 渲染"),
        ("tags", "json", "标签列表", "NULL", "JSON 字符串数组；无标签可为 NULL 或 []"),
        ("created_at", "datetime(3)", "创建时间", "NOT NULL", "UTC；列表按此字段倒序"),
        ("updated_at", "datetime(3)", "更新时间", "NOT NULL", "创建/更新时由服务端写入"),
    ]
    for ri, field in enumerate(fields, 6):
        for ci, v in enumerate(field, 1):
            ws2.cell(ri, ci, v)
        style_body_row(ws2, ri)

    ws2.merge_cells("A14:E14")
    ws2["A14"] = "索引"
    ws2["A14"].font = SECTION_FONT

    for i, h in enumerate(["索引名", "类型", "字段", "说明", "备注"], 1):
        ws2.cell(15, i, h)
    style_header_row(ws2, 15)

    for i, v in enumerate(
        [
            "idx_articles_category_created",
            "INDEX",
            "category, created_at DESC",
            "按分类筛选并按时间倒序",
            "支撑列表接口",
        ],
        1,
    ):
        ws2.cell(16, i, v)
    style_body_row(ws2, 16)

    ws2.merge_cells("A18:E18")
    ws2["A18"] = "枚举与取值"
    ws2["A18"].font = SECTION_FONT

    for i, h in enumerate(["字段", "取值", "含义", "使用场景", "备注"], 1):
        ws2.cell(19, i, h)
    style_header_row(ws2, 19)

    for ri, rowv in enumerate(
        [
            ("category", "有感", "随笔 / 技术笔记", "GET/POST /api/posts", "博客列表主分类"),
            ("category", "诗文", "短诗 / 文本实验", "GET /api/poetry", "诗歌接口"),
        ],
        20,
    ):
        for ci, v in enumerate(rowv, 1):
            ws2.cell(ri, ci, v)
        style_body_row(ws2, ri)

    ws2.row_dimensions[1].height = 28
    ws2.row_dimensions[2].height = 18

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"saved {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
